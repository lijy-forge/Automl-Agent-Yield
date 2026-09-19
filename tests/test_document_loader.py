from __future__ import annotations

from pathlib import Path

import pytest

from yieldmind.document_loader import _strip_repeated_pdf_edges, load_document
from yieldmind.database import YieldMindStore
from yieldmind.knowledge_base import KnowledgeBase, KnowledgeIngestRequest, KnowledgeSearchRequest


def _write_text_pdf(path: Path, text: str) -> None:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode("ascii"))
        payload.extend(obj)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode(
            "ascii"
        )
    )
    path.write_bytes(payload)


def test_load_markdown_preserves_heading_hierarchy(tmp_path: Path) -> None:
    path = tmp_path / "mechanisms.md"
    path.write_text(
        "# Rheology\n\nOverview.\n\n## YODEL\n\nPacking and contact network.\n\n### Limits\n\nCalibrate phi_m.",
        encoding="utf-8",
    )

    loaded = load_document(path)

    assert loaded.title == "Rheology"
    assert [section.section for section in loaded.sections] == [
        "Rheology",
        "Rheology > YODEL",
        "Rheology > YODEL > Limits",
    ]
    assert loaded.source_type == "md"
    assert len(loaded.document_hash) == 64


def test_load_pdf_preserves_page_number(tmp_path: Path) -> None:
    path = tmp_path / "paper.pdf"
    _write_text_pdf(path, "Yield stress and maximum packing fraction")

    loaded = load_document(path, title="Curated paper title")

    assert loaded.title == "Curated paper title"
    assert "maximum packing" in loaded.sections[0].text
    assert loaded.sections[0].page_start == 1
    assert loaded.sections[0].page_end == 1


def test_load_blank_pdf_reports_ocr_requirement(tmp_path: Path) -> None:
    from pypdf import PdfWriter

    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(ValueError, match="require OCR"):
        load_document(path)


def test_pdf_edge_cleanup_only_removes_repeated_headers_and_page_markers() -> None:
    pages = [
        f"Journal Header\nSection {index}\nPage-specific claim {index}\n"
        f"Yield stress is measured.\nConclusion {index}\n{index} of 4"
        for index in range(1, 5)
    ]

    cleaned = _strip_repeated_pdf_edges(pages)

    assert all("Journal Header" not in page for page in cleaned)
    assert all("of 4" not in page for page in cleaned)
    assert all("Yield stress is measured." in page for page in cleaned)
    assert all(f"Page-specific claim {index}" in cleaned[index - 1] for index in range(1, 5))


def test_load_docx_preserves_headings_and_tables(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "protocol.docx"
    document = Document()
    document.add_heading("Measurement", level=1)
    document.add_paragraph("Use a vane geometry for yield stress.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Unit"
    table.rows[0].cells[1].text = "Pa"
    document.save(path)

    loaded = load_document(path)

    assert loaded.title == "Measurement"
    assert loaded.sections[0].section == "Measurement"
    assert "Unit | Pa" in loaded.sections[0].text


def test_load_html_prefers_article_and_preserves_headings(tmp_path: Path) -> None:
    path = tmp_path / "article.html"
    path.write_text(
        "<html><head><title>Fallback</title></head><body><nav>Ignore me</nav>"
        "<article><h1>Yield models</h1><p>Bingham baseline.</p>"
        "<h2>Limits</h2><p>Check wall slip.</p></article></body></html>",
        encoding="utf-8",
    )

    loaded = load_document(path)

    assert loaded.title == "Yield models"
    assert [section.section for section in loaded.sections] == ["Yield models", "Yield models > Limits"]
    assert all("Ignore me" not in section.text for section in loaded.sections)


def test_load_document_rejects_unsupported_type(tmp_path: Path) -> None:
    path = tmp_path / "notes.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported document type"):
        load_document(path)


def test_pdf_ingest_and_search_preserve_provenance(tmp_path: Path) -> None:
    path = tmp_path / "paper.pdf"
    _write_text_pdf(path, "Yield stress depends on packing fraction and contact networks")
    store = YieldMindStore(tmp_path / "yieldmind.sqlite3")
    kb = KnowledgeBase(store=store, chroma_dir=tmp_path / "chroma", collection_name="pdf_provenance")
    metadata = {
        str(path): {
            "title": "Packing paper",
            "source_url": "https://example.org/paper",
            "doi": "https://doi.org/10.1000/example",
            "license": "CC-BY-4.0",
            "metadata": {"publisher": "Test publisher"},
        }
    }

    ingest = kb.ingest(
        KnowledgeIngestRequest(
            paths=[str(path)],
            source_metadata=metadata,
            chunk_size=300,
            chunk_overlap=20,
        )
    )
    document = ingest["documents"][0]
    assert document["status"] == "available"
    assert document["doi"] == "10.1000/example"
    assert ingest["split_version"] == "format_aware_recursive_chars_300_20_v3"

    result = kb.search(
        KnowledgeSearchRequest(query="packing fraction contact networks", retrieval_mode="hybrid", top_k=1)
    )
    hit = result["hits"][0]
    assert hit["page_start"] == 1
    assert hit["page_end"] == 1
    assert hit["source_type"] == "pdf"
    assert hit["source_url"] == "https://example.org/paper"
    assert hit["doi"] == "10.1000/example"
    assert hit["license"] == "CC-BY-4.0"

    listed = kb.list_documents()
    assert listed[0]["metadata"] == {"publisher": "Test publisher"}


def test_doi_identity_prevents_duplicate_after_local_file_move(tmp_path: Path) -> None:
    first = tmp_path / "download-a.pdf"
    second = tmp_path / "download-b.pdf"
    _write_text_pdf(first, "Bingham and Herschel Bulkley yield models")
    second.write_bytes(first.read_bytes())
    store = YieldMindStore(tmp_path / "yieldmind.sqlite3")
    kb = KnowledgeBase(store=store, chroma_dir=tmp_path / "chroma", collection_name="doi_identity")

    first_result = kb.ingest(
        KnowledgeIngestRequest(paths=[str(first)], source_metadata={str(first): {"doi": "10.1000/same"}})
    )
    second_result = kb.ingest(
        KnowledgeIngestRequest(paths=[str(second)], source_metadata={str(second): {"doi": "10.1000/same"}})
    )

    assert first_result["documents"][0]["document_id"] == second_result["documents"][0]["document_id"]
    assert len(kb.list_documents()) == 1
    assert kb.list_documents()[0]["source_path"] == str(second.resolve())


def test_knowledge_search_can_isolate_project_and_literature_corpora(tmp_path: Path) -> None:
    project = tmp_path / "project.md"
    project.write_text("Project retry budget and tool limits.", encoding="utf-8")
    paper = tmp_path / "paper.md"
    paper.write_text("Literature packing fraction and particle contacts.", encoding="utf-8")
    store = YieldMindStore(tmp_path / "yieldmind.sqlite3")
    kb = KnowledgeBase(store=store, chroma_dir=tmp_path / "chroma", collection_name="corpus_filter")
    kb.ingest(KnowledgeIngestRequest(paths=[str(project)]))
    kb.ingest(
        KnowledgeIngestRequest(
            paths=[str(paper)],
            source_metadata={str(paper): {"corpus": "literature"}},
        )
    )

    project_hits = kb.search(
        KnowledgeSearchRequest(query="retry budget", corpus="project", retrieval_mode="bm25")
    )["hits"]
    literature_hits = kb.search(
        KnowledgeSearchRequest(query="packing fraction", corpus="literature", retrieval_mode="bm25")
    )["hits"]
    cross_corpus = kb.search(
        KnowledgeSearchRequest(query="packing fraction", corpus="project", retrieval_mode="bm25")
    )["hits"]

    assert project_hits[0]["corpus"] == "project"
    assert literature_hits[0]["corpus"] == "literature"
    assert cross_corpus == []


def test_vector_search_hides_uncommitted_chroma_chunks(tmp_path: Path) -> None:
    document = tmp_path / "committed.md"
    document.write_text("Committed yield stress evidence.", encoding="utf-8")
    store = YieldMindStore(tmp_path / "yieldmind.sqlite3")
    kb = KnowledgeBase(store=store, chroma_dir=tmp_path / "chroma", collection_name="commit_guard")
    kb.ingest(KnowledgeIngestRequest(paths=[str(document)], chunk_size=300, chunk_overlap=20))

    query = "orphan vector should never be visible"
    collection = kb._collection()
    orphan_ids = [f"orphan_chunk_{index}" for index in range(20)]
    orphan_metadata = [
        {
                "document_id": "missing_document",
                "document_version": "uncommitted",
                "chunk_id": chunk_id,
                "chunk_index": index,
                "title": "Uncommitted",
                "source_path": "/tmp/uncommitted.md",
                "source_type": "md",
                "corpus": "project",
                "source_url": "",
                "doi": "",
                "license": "",
                "section": "",
                "text_hash": "uncommitted",
                "index_version": kb.embedding_profile.index_version,
                "split_version": "uncommitted",
        }
        for index, chunk_id in enumerate(orphan_ids)
    ]
    orphan_vector = kb.embedding.embed_query(query)
    collection.upsert(
        ids=orphan_ids,
        documents=[query] * len(orphan_ids),
        embeddings=[orphan_vector] * len(orphan_ids),
        metadatas=orphan_metadata,
    )

    hits = kb.search(KnowledgeSearchRequest(query=query, retrieval_mode="vector", top_k=5))["hits"]

    assert hits
    assert all(hit["chunk_id"] not in orphan_ids for hit in hits)
    assert all(hit["title"] != "Uncommitted" for hit in hits)
